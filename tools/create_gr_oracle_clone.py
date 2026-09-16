"""Cria cópia temporária e sanitizada das fórmulas GR para recálculo no Excel.

A fonte oficial não é modificada. A saída contém somente planilhas técnicas e
nenhum campo de cliente, obra, local ou orçamento individual.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from audit_maxim_ar_phase3a import sha256

OFFICIAL_SHA = "96514d7818bcbbc1db86f94f86d8ed0239675e9f8d3a6b64df651f30fd90c160"
TECHNICAL_SHEETS = ("GR", "LISTAPERFIS", "LISTAVIDROS", "LISTAFERRA", "LISTADIV", "PFAB")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if sha256(args.source) != OFFICIAL_SHA:
        raise SystemExit("Fonte XLSM não é a versão oficial")
    source = load_workbook(args.source, data_only=False, keep_links=False, keep_vba=False)
    output = Workbook()
    output.remove(output.active)
    copied_formulas = 0
    for sheet_name in TECHNICAL_SHEETS:
        original = source[sheet_name]
        clone = output.create_sheet(sheet_name)
        for row in original:
            for cell in row:
                if cell.value is None:
                    continue
                if sheet_name == "GR":
                    if cell.row == 1 and cell.column in (1, 2, 3, 10):
                        continue
                    if cell.row == 2 and cell.column in (1, 2, 3, 10):
                        continue
                    if cell.column > 26 and cell.row > 200:
                        continue
                clone[cell.coordinate] = cell.value
                copied_formulas += isinstance(cell.value, str) and cell.value.startswith("=")
    # Names used by GR's VLOOKUP formulas. They point only to copied catalogs.
    for name, reference in {
        "LISTAP": "LISTAPERFIS!$A:$F",
        "LISTAV": "LISTAVIDROS!$A:$C",
        "LISTAF": "LISTAFERRA!$A:$C",
        "PAFAB": "PFAB!$A$1:$B$19",
    }.items():
        output.defined_names.add(DefinedName(name, attr_text=reference))
    output.save(args.output)
    print(f"created sanitized GR formula clone: {copied_formulas} formulas")


if __name__ == "__main__":
    main()
